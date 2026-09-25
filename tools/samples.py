"""Four made-up sample files for the in-browser demo. Nothing real is in them.

usage: python tools/samples.py <outdir>
"""
import io
import os
import sys

import msoffcrypto
import openpyxl
from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

TEXT = ("This is a sample document made for the file-checker demo. It contains nothing but "
        "this paragraph, repeated so the file has some body to it. ")


def docx_bytes():
    d = Document()
    d.add_heading("Sample letter", 1)
    for _ in range(3):
        d.add_paragraph(TEXT * 3)
    b = io.BytesIO()
    d.save(b)
    return b.getvalue()


def main():
    out = sys.argv[1]
    os.makedirs(out, exist_ok=True)
    letter = docx_bytes()
    open(os.path.join(out, "letter.docx"), "wb").write(letter)

    enc = io.BytesIO()
    msoffcrypto.OfficeFile(io.BytesIO(letter)).encrypt("sample", enc)
    open(os.path.join(out, "locked.docx"), "wb").write(enc.getvalue())

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Month", "Visitors"])
    for i, m in enumerate(["Jan", "Feb", "Mar", "Apr", "May", "Jun"]):
        ws.append([m, 100 + 17 * i])
    b = io.BytesIO()
    wb.save(b)
    open(os.path.join(out, "report.pdf"), "wb").write(b.getvalue())  # a spreadsheet, misnamed

    b = io.BytesIO()
    c = canvas.Canvas(b, pagesize=A4)
    for page in range(3):
        t = c.beginText(72, 770)
        t.setFont("Helvetica", 11)
        t.textLine(f"Sample PDF, page {page + 1}")
        for _ in range(30):
            t.textLine(TEXT[:90])
        c.drawText(t)
        c.showPage()
    c.save()
    pdf = b.getvalue()
    open(os.path.join(out, "cut.pdf"), "wb").write(pdf[: len(pdf) * 6 // 10])  # cut off mid-file

    for f in sorted(os.listdir(out)):
        print(f, os.path.getsize(os.path.join(out, f)))


if __name__ == "__main__":
    main()
