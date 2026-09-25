"""Build the constructed test set: files whose true type is known because we made them.

Covers what govdocs1 lacks: modern OOXML, encrypted OOXML, ODF/EPUB, email, lying extensions
and truncation. Names are neutral (c0001.ext) so no tool can read the answer from the filename;
the truth is only in manifest.json.

usage: python tools/construct.py <contractnli raw dir> <outdir>
"""
import gzip
import io
import json
import os
import random
import shutil
import sys
import zipfile
from email.message import EmailMessage

import msoffcrypto
import openpyxl
from docx import Document
from PIL import Image
from pptx import Presentation

SEED = 20260925
WORDS = ("the of and to in a is that for it as was with be by on not this are or from at "
         "which but have an they you were his her all one there been if more when will "
         "would who so no river stone lamp window paper garden letter morning road field").split()


class Builder:
    def __init__(self, outdir, rng):
        self.outdir, self.rng, self.n, self.manifest = outdir, rng, 0, []

    def path(self, ext):
        self.n += 1
        name = f"c{self.n:04d}.{ext}"
        return name, os.path.join(self.outdir, name)

    def add(self, name, label, origin, flags=()):
        self.manifest.append({"file": name, "label": label, "flags": sorted(flags), "origin": origin})

    def sentence(self, n=12):
        return " ".join(self.rng.choice(WORDS) for _ in range(n)).capitalize() + "."


def make_docx(b, path):
    d = Document()
    d.add_heading(b.sentence(4), 1)
    for _ in range(b.rng.randint(2, 8)):
        d.add_paragraph(b.sentence(b.rng.randint(8, 30)))
    d.save(path)


def make_xlsx(b, path):
    wb = openpyxl.Workbook()
    ws = wb.active
    for r in range(b.rng.randint(5, 40)):
        ws.append([b.rng.randint(0, 999) if c % 2 else b.rng.choice(WORDS) for c in range(6)])
    wb.save(path)


def make_pptx(b, path):
    p = Presentation()
    for _ in range(b.rng.randint(1, 5)):
        s = p.slides.add_slide(p.slide_layouts[1])
        s.shapes.title.text = b.sentence(4)
        s.placeholders[1].text = b.sentence(15)
    p.save(path)


def make_odf(b, path, kind):
    mime = {"odt": "application/vnd.oasis.opendocument.text",
            "ods": "application/vnd.oasis.opendocument.spreadsheet",
            "odp": "application/vnd.oasis.opendocument.presentation",
            "epub": "application/epub+zip"}[kind]
    with zipfile.ZipFile(path, "w") as z:
        # ODF and EPUB both require the stored, uncompressed "mimetype" entry to come first.
        z.writestr(zipfile.ZipInfo("mimetype"), mime, compress_type=zipfile.ZIP_STORED)
        if kind == "epub":
            z.writestr("META-INF/container.xml", '<?xml version="1.0"?><container version="1.0" '
                       'xmlns="urn:oasis:names:tc:opendocument:xmlns:container"><rootfiles>'
                       '<rootfile full-path="content.opf" media-type="application/oebps-package+xml"/>'
                       '</rootfiles></container>', compress_type=zipfile.ZIP_DEFLATED)
            z.writestr("content.opf", "<package/>", compress_type=zipfile.ZIP_DEFLATED)
            z.writestr("ch1.xhtml", f"<html><body><p>{b.sentence(40)}</p></body></html>",
                       compress_type=zipfile.ZIP_DEFLATED)
        else:
            z.writestr("META-INF/manifest.xml", f'<?xml version="1.0"?><manifest:manifest '
                       f'xmlns:manifest="urn:oasis:names:tc:opendocument:xmlns:manifest:1.0">'
                       f'<manifest:file-entry manifest:full-path="/" manifest:media-type="{mime}"/>'
                       f'</manifest:manifest>', compress_type=zipfile.ZIP_DEFLATED)
            z.writestr("content.xml", f"<office:document-content><p>{b.sentence(40)}</p>"
                       f"</office:document-content>", compress_type=zipfile.ZIP_DEFLATED)


def make_image(b, path, fmt):
    w, h = b.rng.randint(16, 400), b.rng.randint(16, 400)
    img = Image.effect_noise((w, h), b.rng.randint(10, 100)).convert("RGB")
    img.save(path, format=fmt)


def text_file(path, s):
    with open(path, "w", encoding="utf-8", newline="\n") as f:
        f.write(s)


def main():
    raw, outdir = sys.argv[1:3]
    rng = random.Random(SEED)
    if os.path.isdir(outdir):
        shutil.rmtree(outdir)
    os.makedirs(outdir)
    b = Builder(outdir, rng)
    made = {"docx": [], "xlsx": [], "pptx": [], "pdf": [], "png": []}

    # Modern OOXML, freshly written.
    for kind, fn in (("docx", make_docx), ("xlsx", make_xlsx), ("pptx", make_pptx)):
        for _ in range(20):
            name, p = b.path(kind)
            fn(b, p)
            b.add(name, kind, "python-" + kind)
            made[kind].append(p)

    # Encrypted OOXML: password-protected files are OLE containers holding an EncryptedPackage.
    for kind in ("docx", "xlsx", "pptx"):
        for src in made[kind][:10]:
            name, p = b.path(kind)
            with open(src, "rb") as f, open(p, "wb") as out:
                msoffcrypto.OfficeFile(f).encrypt("correct horse", out)
            b.add(name, "office-encrypted", "msoffcrypto-" + kind, {"encrypted"})

    # Real PDFs from ContractNLI (CC BY 4.0), already on disk.
    pdfs = sorted(f for f in os.listdir(raw) if f.lower().endswith(".pdf"))
    for f in rng.sample(pdfs, 20):
        name, p = b.path("pdf")
        shutil.copyfile(os.path.join(raw, f), p)
        b.add(name, "pdf", "contractnli")
        made["pdf"].append(p)

    # ODF and EPUB, hand-assembled to spec (no office suite here).
    for kind in ("odt", "ods", "odp", "epub"):
        for _ in range(5):
            name, p = b.path(kind)
            make_odf(b, p, kind)
            b.add(name, kind, "zipfile-" + kind)

    # Images.
    for fmt, ext in (("PNG", "png"), ("JPEG", "jpg"), ("GIF", "gif"), ("TIFF", "tif"),
                     ("BMP", "bmp"), ("WEBP", "webp")):
        for _ in range(5):
            name, p = b.path(ext)
            make_image(b, p, fmt)
            b.add(name, {"jpg": "jpeg", "tif": "tiff"}.get(ext, ext), "pillow")
            if ext == "png":
                made["png"].append(p)

    # Archives: a plain zip must not be mistaken for OOXML, and vice versa.
    for _ in range(5):
        name, p = b.path("zip")
        with zipfile.ZipFile(p, "w", zipfile.ZIP_DEFLATED) as z:
            for i in range(rng.randint(1, 4)):
                z.writestr(f"notes/{i}.txt", b.sentence(50))
        b.add(name, "zip", "zipfile")
        name, p = b.path("gz")
        with gzip.open(p, "wb") as g:
            g.write(b.sentence(200).encode())
        b.add(name, "gzip", "gzip")

    # Text-ish, labelled by construction (these are easy cases; reported separately).
    for _ in range(10):
        name, p = b.path("json")
        text_file(p, json.dumps({"id": rng.randint(1, 9999), "tags": rng.sample(WORDS, 3),
                                 "note": b.sentence()}, indent=rng.choice([None, 2])))
        b.add(name, "json", "construct")
        name, p = b.path("xml")
        decl = '<?xml version="1.0" encoding="UTF-8"?>\n' if rng.random() < 0.5 else ""
        text_file(p, decl + "<records>" + "".join(f"<r n=\"{i}\">{b.sentence(5)}</r>"
                                                   for i in range(rng.randint(2, 20))) + "</records>\n")
        b.add(name, "xml", "construct")
        name, p = b.path("html")
        doctype = "<!DOCTYPE html>\n" if rng.random() < 0.5 else ""
        text_file(p, f"{doctype}<html><head><title>{b.sentence(3)}</title></head><body>"
                     f"<p>{b.sentence(30)}</p></body></html>\n")
        b.add(name, "html", "construct")
        name, p = b.path("csv")
        rows = ["id,name,amount,city"] + [f"{i},{rng.choice(WORDS)},{rng.randint(1, 900)}.{rng.randint(0, 99):02d},"
                                         f"{rng.choice(WORDS)}" for i in range(rng.randint(3, 60))]
        text_file(p, "\n".join(rows) + "\n")
        b.add(name, "csv", "construct")
        name, p = b.path("txt")
        text_file(p, "\n\n".join(b.sentence(rng.randint(10, 40)) for _ in range(rng.randint(1, 6))) + "\n")
        b.add(name, "text", "construct")
        name, p = b.path("eml")
        m = EmailMessage()
        m["From"], m["To"] = "a@example.org", "b@example.org"
        m["Subject"], m["Date"] = b.sentence(4), "Thu, 24 Sep 2026 10:00:00 +0000"
        m.set_content(b.sentence(40))
        if rng.random() < 0.5:
            m.add_attachment(b.sentence(20).encode(), maintype="application", subtype="octet-stream",
                             filename="a.bin")
        with open(p, "wb") as f:
            f.write(bytes(m))
        b.add(name, "eml", "email")
    for _ in range(5):
        name, p = b.path("rtf")
        text_file(p, "{\\rtf1\\ansi\\deff0 {\\fonttbl {\\f0 Times;}}\\f0\\fs24 " + b.sentence(30) + "\\par}\n")
        b.add(name, "rtf", "construct")

    # Lying extensions: same bytes, wrong name. The label is the true type.
    lies = (("docx", "pdf"), ("pdf", "docx"), ("xlsx", "csv"), ("png", "jpg"), ("pptx", "zip"))
    for src_kind, fake_ext in lies:
        for src in made[src_kind][-4:]:
            name, p = b.path(fake_ext)
            shutil.copyfile(src, p)
            b.add(name, src_kind, "renamed", {"extension-mismatch"})

    # Truncated: the first half of a real file. The label stays the original type.
    for kind in ("pdf", "docx", "xlsx"):
        for src in made[kind][:5]:
            name, p = b.path(kind)
            data = open(src, "rb").read()
            with open(p, "wb") as f:
                f.write(data[: len(data) // 2])
            b.add(name, kind, "truncated", {"truncated"})

    with open(os.path.join(outdir, "manifest.json"), "w", encoding="utf-8") as f:
        json.dump(b.manifest, f, indent=1)
    counts = {}
    for e in b.manifest:
        counts[e["label"]] = counts.get(e["label"], 0) + 1
    print(len(b.manifest), "files", json.dumps(counts, sort_keys=True))


if __name__ == "__main__":
    main()
